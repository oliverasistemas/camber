#!/usr/bin/env python3
"""Bootstrap pipeline inputs from Brad's monthly collections workbook.

The workbook ("Camber Collections Report - YYYY.MM.DD (vFormulas).xlsx") embeds
copies of the raw Vervent sFTP reports plus Camber's own reports, manual
overrides, and the computed monthly tape. This script splits it into the files
the pipeline consumes, so a month can be replayed end-to-end without sFTP
access, and emits Brad's computed columns as the expected tie-out values.

Outputs (under pipeline/data/):
  inbox/<asof>/vervent_trans_history.csv     raw servicer transaction history
  inbox/<asof>/vervent_asset_summary.csv     raw servicer asset summary
  inbox/<asof>/vervent_contract_summary.csv  raw servicer contract summary
  inbox/<asof>/vervent_eot_charges.csv       raw servicer end-of-term charges
  inbox/<asof>/camber_lease_term_report.csv  Camber lease termination report
  inbox/<asof>/cb_offboarded_accounts.csv    CB offboarded account list
  seed/pool_2026_1.csv                       static lease/vehicle attributes
  seed/prior_tape_2026-07-31.csv             prior-month tape carry-forward
  overrides/overrides_<asof>.csv             manual overrides for the month
  expected/expected_tape_<asof>.csv          Brad's computed values (tie-out)
"""

import csv
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

ASOF = "2026-08-31"
HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
DEFAULT_WORKBOOK = (
    HERE.parent
    / "backlog"
    / "collections"
    / "Camber Collections Report - 2026.08.31 (vFormulas1.5).xlsx"
)


def norm_col(name):
    """Normalize a header to snake_case (shared convention with the DDL)."""
    s = str(name).strip().lower()
    s = re.sub(r"[^a-z0-9]+", "_", s).strip("_")
    if re.match(r"^[0-9]", s):
        s = "c_" + s
    return s


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v)


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    print(f"wrote {path.relative_to(HERE)} ({len(rows)} rows)")


def extract_sheet(wb, sheet, out_path, header_row=2, max_col=None):
    """Dump a sheet verbatim (normalized headers, ISO dates) to CSV."""
    ws = wb[sheet]
    it = ws.iter_rows(min_row=header_row, max_col=max_col, values_only=True)
    raw_header = next(it)
    # trim trailing all-None header cells
    while raw_header and raw_header[-1] is None:
        raw_header = raw_header[:-1]
    ncol = len(raw_header)
    header = [norm_col(h) if h is not None else f"col_{i+1}" for i, h in enumerate(raw_header)]
    rows = []
    for r in it:
        vals = [fmt(v) for v in r[:ncol]]
        if any(v != "" for v in vals):
            rows.append(vals)
    write_csv(out_path, header, rows)


# Column letters in '2026-1 Aug (Formulas)' (see Field Dictionary sheet)
def col_idx(letter):
    return column_index_from_string(letter) - 1


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    print(f"loading {path} ...")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    inbox = DATA / "inbox" / ASOF
    extract_sheet(wb, "Aug 1-31 TransHist", inbox / "vervent_trans_history.csv")
    extract_sheet(wb, "Aug 31 AssetSummary", inbox / "vervent_asset_summary.csv")
    extract_sheet(wb, "Aug 31 ContractSummary", inbox / "vervent_contract_summary.csv")
    extract_sheet(wb, "Vervent - EOT Charges", inbox / "vervent_eot_charges.csv", max_col=10)
    extract_sheet(wb, "Camber Lease Term Report", inbox / "camber_lease_term_report.csv")
    extract_sheet(wb, "CB Offboarded Accounts", inbox / "cb_offboarded_accounts.csv", header_row=1)

    ws = wb["2026-1 Aug (Formulas)"]
    tape_rows = list(ws.iter_rows(min_row=3, max_row=2230, values_only=True))
    tape_rows = [r for r in tape_rows if r[0] is not None]

    # --- static pool attributes (columns A..Y) --------------------------------
    pool_cols = [
        ("A", "lease_id"), ("B", "vin"), ("C", "model_year"), ("D", "make"),
        ("E", "yr_make_model"), ("F", "fuel_type"), ("G", "new_used"),
        ("H", "open_closed_lease"), ("I", "business_personal"),
        ("J", "lease_extension"), ("K", "lessee_fico"), ("L", "lessee_state"),
        ("M", "lease_term_months"), ("N", "lease_date"), ("O", "maturity_date"),
        ("P", "orig_cost"), ("Q", "contract_residual_value"), ("R", "pmt_reg"),
        ("S", "pmt_freq"), ("T", "alg_model_used"), ("U", "lease_dcf"),
        ("V", "par_value_lessor_yield_dcf_20251031"),
        ("W", "par_value_9pct_payments_20251031"),
        ("X", "par_value_9pct_residual_20251031"),
        ("Y", "par_value_9pct_dcf_20251031"),
    ]
    write_csv(
        DATA / "seed" / "pool_2026_1.csv",
        [n for _, n in pool_cols],
        [[fmt(r[col_idx(c)]) for c, _ in pool_cols] for r in tape_rows],
    )

    # --- prior-month (2026-07-31) tape carry-forward --------------------------
    prior_cols = [
        ("A", "lease_id"), ("AA", "camber_lease_status"), ("AK", "sold_month"),
        ("AN", "net_sales_proceeds"), ("BJ", "total_ar_final"),
        ("BT", "pv9_total_excl_ppt"), ("BY", "cum_cash_flow"),
        ("BZ", "cum_property_tax"), ("CA", "cum_sales_tax"),
    ]
    write_csv(
        DATA / "seed" / "prior_tape_2026-07-31.csv",
        [n for _, n in prior_cols],
        [[fmt(r[col_idx(c)]) for c, _ in prior_cols] for r in tape_rows],
    )

    # --- manual overrides for the current month -------------------------------
    ov_specs = [
        ("CU", "sold_month", "Camber sold-date override (sale confirmed outside servicer data)"),
        ("CX", "sale_proceeds", "Camber sale-proceeds override (e.g. Copart settlement differs from Vervent)"),
        ("DG", "past_due_receivables", "Camber past-due receivables override"),
        ("EA", "days_delinquent", "Camber days-delinquent override"),
        ("DW", "eol_ar_exclusion", "Exclude end-of-lease receivable from total AR"),
        ("CK", "under_review", "Lease under manual review"),
    ]
    ov_rows = []
    for r in tape_rows:
        lease_id = fmt(r[0])
        for letter, field, reason in ov_specs:
            v = r[col_idx(letter)]
            if v in (None, ""):
                continue
            if field == "under_review" and v == "N":
                continue
            ov_rows.append([lease_id, ASOF, field, fmt(v), "brad_workbook", reason])

    # Hidden overrides: computed cells where Brad replaced the formula with a
    # literal value, or edited the row's formula away from the column's norm.
    # Emit them as explicit overrides so the pipeline reproduces his tape.
    scan_cols = {
        "CL": "lease_status",
        "CZ": "period_cash_flow",
        "DD": "period_payment_excl_taxes",
        "CV": "sold_month",
        "CY": "sale_proceeds",
        "DH": "past_due_receivables",
        "EB": "days_delinquent",
    }
    from openpyxl.worksheet.formula import ArrayFormula

    print("scanning for hardcoded / edited formula cells ...")
    wbf = openpyxl.load_workbook(path, read_only=True, data_only=False)
    wsf = wbf["2026-1 Aug (Formulas)"]
    frows = [r for r in wsf.iter_rows(min_row=3, max_row=2230)
             if r[0].value is not None]
    data_by_id = {fmt(r[0]): r for r in tape_rows}
    for letter, field in scan_cols.items():
        ci = col_idx(letter)
        texts = []
        for fr in frows:
            v = fr[ci].value
            texts.append(v.text if isinstance(v, ArrayFormula) else v)
        formulas = [t for t in texts if isinstance(t, str) and t.startswith("=")]
        modal = max(set(formulas), key=formulas.count) if formulas else None
        for fr, t in zip(frows, texts):
            lease_id = fmt(fr[0].value)
            if t is None or (isinstance(t, str) and t.startswith("=") and t == modal):
                continue
            cached = data_by_id[lease_id][ci]
            kind = ("row formula edited away from column formula"
                    if isinstance(t, str) and t.startswith("=")
                    else "formula replaced with a literal value")
            ov_rows.append([lease_id, ASOF, field, fmt(cached),
                            "brad_workbook_cell", f"{kind} (col {letter})"])
            print(f"  {lease_id} {letter}/{field}: {kind} -> {fmt(cached)}")
    write_csv(
        DATA / "overrides" / f"overrides_{ASOF}.csv",
        ["lease_id", "asof_date", "field", "value", "source", "reason"],
        ov_rows,
    )

    # --- expected (Brad-computed) values for tie-out ---------------------------
    exp_cols = [
        ("A", "lease_id"), ("CL", "camber_lease_status"), ("CM", "vervent_lease_status"),
        ("CN", "vervent_asset_status"), ("CO", "termination_type"),
        ("CV", "sold_month"), ("CW", "vervent_net_sale_proceeds"),
        ("CY", "net_sales_proceeds"), ("CZ", "period_total_cash_flow"),
        ("DA", "period_property_tax"), ("DB", "period_sales_tax"),
        ("DD", "period_payment_excl_prop_sales_tax"),
        ("DH", "past_due_receivables"), ("DU", "eol_ar_final"),
        ("DX", "total_ar_final"), ("EB", "days_delinquent"),
        ("ED", "n_pmts_remaining"), ("EE", "pv_ly_payments"),
        ("EF", "pv_ly_residual"), ("EG", "pv_ly_total"),
        ("EH", "pv9_payments"), ("EI", "pv9_residual"),
        ("EJ", "pv9_total_excl_ppt"), ("EK", "ppt_proration"), ("EL", "pv9_total"),
        ("EM", "eligible_flag"), ("EN", "pv9_eligible"),
        ("EO", "cum_cash_flow"), ("ET", "cum_payment_final"),
        ("EZ", "cb_offboarded"),
    ]
    write_csv(
        DATA / "expected" / f"expected_tape_{ASOF}.csv",
        [n for _, n in exp_cols],
        [[fmt(r[col_idx(c)]) for c, _ in exp_cols] for r in tape_rows],
    )

    print("done.")


if __name__ == "__main__":
    main()
